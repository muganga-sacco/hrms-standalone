# Copyright (c) 2026, Frappe Technologies Pvt. Ltd. and contributors
# License: GNU General Public License v3. See license.txt

from __future__ import annotations

import io
import os
import re
from copy import deepcopy
from datetime import date

import frappe
from frappe import _
from frappe.utils import getdate

from hrms.payroll.payroll_import.payroll_month_utils import (
	normalize_employee_id_number,
	normalize_payroll_month,
	payroll_period_label_for_month,
)


def get_template_path() -> str:
	return frappe.get_app_path("hrms", "payroll", "templates", "payslip_template.docx")


def _payslip_get_all(*args, **kwargs):
	"""Payroll rows for payslip generation (caller must enforce access on the API)."""
	kwargs["ignore_permissions"] = True
	return frappe.get_all(*args, **kwargs)


def _clear_paragraph_runs(paragraph) -> None:
	for run in list(paragraph.runs):
		run._element.getparent().remove(run._element)


def _remove_drawings(element) -> None:
	from docx.oxml.ns import qn

	drawing_tag = qn("w:drawing")
	for drawing in list(element.iter(drawing_tag)):
		parent = drawing.getparent()
		if parent is not None:
			parent.remove(drawing)


def _clear_template_logos(doc) -> None:
	"""Template ships logos in the section header (often duplicated) and sometimes the body."""
	for section in doc.sections:
		for paragraph in section.header.paragraphs:
			_clear_paragraph_runs(paragraph)
			_remove_drawings(paragraph._element)
		_remove_drawings(section.header._element)
		for paragraph in section.footer.paragraphs:
			_clear_paragraph_runs(paragraph)
			_remove_drawings(paragraph._element)

	for paragraph in doc.paragraphs:
		text = (paragraph.text or "").strip().upper()
		if "PAYSLIP" in text or text.startswith("EMPLOYEE") or text.startswith("POSITION"):
			break
		_clear_paragraph_runs(paragraph)
		_remove_drawings(paragraph._element)


def _inject_payslip_logo(doc, record: dict) -> None:
	from hrms.payroll.payroll_import.payroll_sheet_layout import (
		get_company_logo_path,
		get_payroll_excel_upload_for_month,
	)

	company = record.get("company")
	payroll_excel_upload = record.get("payroll_excel_upload") or get_payroll_excel_upload_for_month(
		company, record.get("payroll_month")
	)
	logo_path = get_company_logo_path(company, payroll_excel_upload)
	try:
		from docx.shared import Inches

		_clear_template_logos(doc)
		if not logo_path:
			return
		paragraph = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
		_clear_paragraph_runs(paragraph)
		_remove_drawings(paragraph._element)
		run = paragraph.add_run()
		run.add_picture(logo_path, height=Inches(0.85))
	except Exception:
		frappe.log_error(title="Payslip logo")


def _site_file_path(path: str | None) -> str | None:
	if not path:
		return None
	if os.path.isabs(path) and os.path.exists(path):
		return path
	rooted = os.path.abspath(os.path.join(frappe.utils.get_bench_path(), "sites", str(path).lstrip("./")))
	if os.path.exists(rooted):
		return rooted
	if os.path.exists(path):
		return os.path.abspath(path)
	return None


def _employee_photo_path(record: dict, employee_id_number: str | None = None) -> str | None:
	"""Photo stored on the Employee, or on the linked User when the employee image is empty."""
	from hrms.payroll.payroll_import.payroll_employee_link import find_employee_for_payroll_id
	from hrms.payroll.payroll_import.payroll_signoff import signature_image_path

	employee = (record.get("employee") or "").strip()
	if employee and not frappe.db.exists("Employee", employee):
		employee = ""
	if not employee:
		employee = (
			find_employee_for_payroll_id(record.get("employee_id_number") or employee_id_number) or ""
		)
	if not employee:
		return None

	image = frappe.db.get_value("Employee", employee, "image")
	if not image:
		user_id = frappe.db.get_value("Employee", employee, "user_id")
		if user_id:
			image = frappe.db.get_value("User", user_id, "user_image")
	return _site_file_path(signature_image_path(image))


def _payslip_logo_path(record: dict) -> str | None:
	from hrms.payroll.payroll_import.payroll_sheet_layout import (
		get_company_logo_path,
		get_payroll_excel_upload_for_month,
	)

	company = record.get("company")
	payroll_excel_upload = record.get("payroll_excel_upload") or get_payroll_excel_upload_for_month(
		company, record.get("payroll_month")
	)
	return _site_file_path(get_company_logo_path(company, payroll_excel_upload))


def _remove_table_borders(table) -> None:
	from docx.oxml import OxmlElement
	from docx.oxml.ns import qn

	tbl_pr = table._tbl.tblPr
	existing = tbl_pr.find(qn("w:tblBorders"))
	if existing is not None:
		tbl_pr.remove(existing)
	borders = OxmlElement("w:tblBorders")
	for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
		element = OxmlElement(f"w:{edge}")
		element.set(qn("w:val"), "nil")
		element.set(qn("w:sz"), "0")
		element.set(qn("w:space"), "0")
		element.set(qn("w:color"), "auto")
		borders.append(element)
	tbl_pr.append(borders)


def _set_full_width_table(table, width_twips: int) -> None:
	from docx.oxml import OxmlElement
	from docx.oxml.ns import qn

	tbl = table._tbl
	tbl_pr = tbl.tblPr
	tbl_w = tbl_pr.find(qn("w:tblW"))
	if tbl_w is None:
		tbl_w = OxmlElement("w:tblW")
		tbl_pr.append(tbl_w)
	tbl_w.set(qn("w:w"), str(width_twips))
	tbl_w.set(qn("w:type"), "dxa")

	layout = tbl_pr.find(qn("w:tblLayout"))
	if layout is None:
		layout = OxmlElement("w:tblLayout")
		tbl_pr.append(layout)
	layout.set(qn("w:type"), "fixed")

	indent = tbl_pr.find(qn("w:tblInd"))
	if indent is None:
		indent = OxmlElement("w:tblInd")
		tbl_pr.append(indent)
	indent.set(qn("w:w"), "0")
	indent.set(qn("w:type"), "dxa")

	grid = tbl.find(qn("w:tblGrid"))
	cols = list(grid.findall(qn("w:gridCol"))) if grid is not None else []
	left = width_twips // 2
	right = width_twips - left
	widths = (left, right)
	for col, width in zip(cols, widths, strict=False):
		col.set(qn("w:w"), str(width))
	for cell, width in zip(table.rows[0].cells, widths, strict=False):
		tc_pr = cell._tc.get_or_add_tcPr()
		tc_w = tc_pr.find(qn("w:tcW"))
		if tc_w is None:
			tc_w = OxmlElement("w:tcW")
			tc_pr.append(tc_w)
		tc_w.set(qn("w:w"), str(width))
		tc_w.set(qn("w:type"), "dxa")
		v_align = tc_pr.find(qn("w:vAlign"))
		if v_align is None:
			v_align = OxmlElement("w:vAlign")
			tc_pr.append(v_align)
		v_align.set(qn("w:val"), "center")


def _add_header_picture(cell, path: str, align, *, logo: bool) -> None:
	from docx.shared import Inches, Pt
	from PIL import Image

	paragraph = cell.paragraphs[0]
	paragraph.alignment = align
	paragraph.paragraph_format.space_before = Pt(0)
	paragraph.paragraph_format.space_after = Pt(0)
	run = paragraph.add_run()
	if logo:
		run.add_picture(path, height=Inches(0.85))
		return
	with Image.open(path) as image:
		wide = image.width >= image.height
	if wide:
		run.add_picture(path, width=Inches(0.95))
	else:
		run.add_picture(path, height=Inches(0.95))


def _inject_employee_photo(doc, record: dict, employee_id_number: str | None = None) -> None:
	"""Logo on the left, employee photo in the top-right, PAYSLIP centered below."""
	from docx.enum.text import WD_ALIGN_PARAGRAPH
	from docx.shared import Emu

	photo_path = _employee_photo_path(record, employee_id_number)
	if not photo_path or not doc.paragraphs:
		return
	logo_path = _payslip_logo_path(record)
	logo_para = doc.paragraphs[0]
	try:
		section = doc.sections[0]
		usable = int(section.page_width - section.left_margin - section.right_margin)
		header = doc.add_table(rows=1, cols=2)
		_remove_table_borders(header)
		_set_full_width_table(header, int(Emu(usable).twips))
		if logo_path:
			_add_header_picture(header.rows[0].cells[0], logo_path, WD_ALIGN_PARAGRAPH.LEFT, logo=True)
		_add_header_picture(header.rows[0].cells[1], photo_path, WD_ALIGN_PARAGRAPH.RIGHT, logo=False)
		logo_para._p.addprevious(header._tbl)
		if not (logo_para.text or "").strip().upper().startswith("PAYSLIP"):
			logo_para._element.getparent().remove(logo_para._element)
	except Exception:
		tbl = header._tbl if "header" in locals() else None
		if tbl is not None and tbl.getparent() is not None:
			tbl.getparent().remove(tbl)
		frappe.log_error(title="Payslip employee photo")


def _clear_template_signatories(doc) -> None:
	"""Drop the fixed names and the blank lines under 'For MUGANGA SACCO'."""
	seen_for = False
	remove = []
	for paragraph in doc.paragraphs:
		text = (paragraph.text or "").strip()
		if text.upper().startswith("FOR "):
			seen_for = True
			continue
		if seen_for:
			remove.append(paragraph._element)
	for element in remove:
		parent = element.getparent()
		if parent is not None:
			parent.remove(element)


def _append_payslip_signoff(doc, signoff: dict | None) -> None:
	"""Verified by (HR) and Approved by (DAF) for whoever approved this request."""
	if not signoff:
		return
	blocks = [block for block in (signoff.get("verified_by"), signoff.get("approved_by")) if block]
	if not blocks:
		return
	_clear_template_signatories(doc)
	try:
		from docx.oxml import OxmlElement
		from docx.oxml.ns import qn
		from docx.shared import Emu, Inches, Pt

		table = doc.add_table(rows=1, cols=len(blocks))
		_remove_table_borders(table)
		if len(blocks) == 2:
			section = doc.sections[0]
			usable = int(section.page_width - section.left_margin - section.right_margin)
			_set_full_width_table(table, int(Emu(usable).twips))
		tr_pr = table.rows[0]._tr.get_or_add_trPr()
		tr_pr.append(OxmlElement("w:cantSplit"))

		for cell, block in zip(table.rows[0].cells, blocks, strict=False):
			tc_pr = cell._tc.get_or_add_tcPr()
			v_align = tc_pr.find(qn("w:vAlign"))
			if v_align is None:
				v_align = OxmlElement("w:vAlign")
				tc_pr.append(v_align)
			v_align.set(qn("w:val"), "top")

			def _tight(paragraph, before=0):
				paragraph.paragraph_format.space_before = Pt(before)
				paragraph.paragraph_format.space_after = Pt(0)
				paragraph.paragraph_format.line_spacing = 1.0

			heading = cell.paragraphs[0]
			_tight(heading, 2)
			heading_run = heading.add_run(block.get("heading") or "")
			heading_run.bold = True
			heading_run.font.size = Pt(11)

			name = cell.add_paragraph()
			_tight(name)
			name_run = name.add_run(block.get("name") or "")
			name_run.bold = True
			name_run.font.size = Pt(11)

			if block.get("title"):
				title = cell.add_paragraph()
				_tight(title)
				title_run = title.add_run(block.get("title"))
				title_run.font.size = Pt(9)

			image_path = block.get("image_path")
			if image_path and os.path.exists(image_path):
				signature = cell.add_paragraph()
				_tight(signature)
				signature.add_run().add_picture(image_path, height=Inches(0.32))

			if block.get("approved_on"):
				when = cell.add_paragraph()
				_tight(when)
				when_run = when.add_run(_("Approved on {0}").format(block.get("approved_on")))
				when_run.font.size = Pt(8)
	except Exception:
		frappe.log_error(title="Payslip sign-off")


PAYSLIP_FIELDS = [
	"name",
	"modified",
	"employee",
	"company",
	"payroll_excel_upload",
	"employee_id_number",
	"employee_name",
	"position",
	"department",
	"payroll_period_label",
	"payroll_month",
	"payment_date",
	"basic_salary",
	"transport_allowance",
	"ict_allowance",
	"itc_allowance",
	"gross_salary",
	"payroll_sheet_type",
	"paye",
	"pension_employee",
	"maternity_employee",
	"maternity_employer",
	"cbhi",
	"rpf",
	"insurance_sanlam",
	"insurance_prime",
	"muganga_sacco_social_fund",
	"brd_loan",
	"sport_subscription",
	"other_deductions",
	"total_deductions",
	"net_salary",
	"net_paid",
	"take_home",
	"account_number",
	"bank",
]

MAX_PAYSLIP_MONTHS = 24
MAX_PAYSLIP_ROWS_ONE_PAGE = 3
# Template grid is 14 columns, including ICT Allowance. Transport, TC Allowance,
# and Bonus are inserted. ICT is removed when no row has an ICT amount.
PAYSLIP_TABLE_COLS_WITH_ICT = 17
PAYSLIP_TABLE_COLS_WITHOUT_ICT = 16
BONUS_COLUMN_WIDTH = 1600
ALLOWANCE_COLUMN_WIDTH = 1154
TRANSPORT_COLUMN_WIDTH = 2200


def get_payroll_records_for_payslip(
	employee_id_number: str,
	months: int | None = None,
	from_month: str | None = None,
	to_month: str | None = None,
	payroll_month: str | None = None,
) -> list[dict]:
	employee_id_number = normalize_employee_id_number(employee_id_number)
	if not employee_id_number:
		return []

	scope = _resolve_payslip_employee_scope(employee_id_number)
	limit = None
	month_filters: dict = {}

	if payroll_month:
		month_filters["payroll_month"] = normalize_payroll_month(payroll_month)
	elif from_month and to_month:
		month_filters["payroll_month"] = [
			"between",
			[normalize_payroll_month(from_month), normalize_payroll_month(to_month)],
		]
	elif from_month:
		month_filters["payroll_month"] = [">=", normalize_payroll_month(from_month)]
	elif to_month:
		month_filters["payroll_month"] = ["<=", normalize_payroll_month(to_month)]
	else:
		limit = min(int(months or 3), MAX_PAYSLIP_MONTHS)

	fetch_limit = (limit * 4) if limit else None
	records = _fetch_imported_records_for_scope(scope, month_filters, fetch_limit)
	merged = _merge_records_by_payroll_month(records, scope)
	if limit:
		merged = merged[:limit]
	return merged


def _normalize_person_name(name) -> str:
	return re.sub(r"\s+", " ", (name or "").strip()).upper()


def _is_grand_total_record(record: dict) -> bool:
	for field in ("employee_id_number", "employee_name"):
		val = (record.get(field) or "").strip().upper()
		if val.startswith("GRAND TOTAL"):
			return True
	return False


def _is_lumpsum_record(record: dict) -> bool:
	return (record.get("payroll_sheet_type") or "Staff Payroll").strip() == "Lumpsum"


def _is_variable_record(record: dict) -> bool:
	return (record.get("payroll_sheet_type") or "").strip() == "Variable"


def _add_employee_id_variant(id_variants: set[str], val) -> None:
	text = (val or "").strip()
	if not text or text.upper().startswith("GRAND TOTAL"):
		return
	id_variants.add(text)


def _resolve_payslip_employee_scope(employee_id_number: str) -> dict:
	from hrms.payroll.payroll_import.payroll_employee_link import (
		find_employee_for_payroll_id,
		payroll_id_variants,
	)

	employee_id_number = (employee_id_number or "").strip()
	employee = find_employee_for_payroll_id(employee_id_number)

	id_variants: set[str] = set()
	name_keys: set[str] = set()
	account_numbers: set[str] = set()

	for variant in payroll_id_variants(employee_id_number):
		_add_employee_id_variant(id_variants, variant)
	_add_employee_id_variant(id_variants, employee_id_number)
	if employee:
		_add_employee_id_variant(
			id_variants, frappe.db.get_value("Employee", employee, "employee_number")
		)
		emp_name = frappe.db.get_value("Employee", employee, "employee_name")
		if emp_name:
			name_keys.add(_normalize_person_name(emp_name))

	or_filters: list[list] = []
	if employee:
		or_filters.append(["employee", "=", employee])
	for eid in list(id_variants):
		or_filters.append(["employee_id_number", "=", eid])

	if or_filters:
		refs = _payslip_get_all(
			"Imported Payroll Record",
			filters={"docstatus": 1},
			or_filters=or_filters,
			fields=["employee", "employee_id_number", "employee_name", "account_number"],
			limit=500,
		)
		for row in refs:
			if row.get("employee") and not employee:
				employee = row.get("employee")
			_add_employee_id_variant(id_variants, row.get("employee_id_number"))
			name_key = _normalize_person_name(row.get("employee_name"))
			if name_key:
				name_keys.add(name_key)
			account = (row.get("account_number") or "").strip()
			if account:
				account_numbers.add(account)

	return {
		"employee": employee,
		"id_variants": id_variants,
		"id_variants_upper": {v.upper() for v in id_variants},
		"name_keys": name_keys,
		"account_numbers": account_numbers,
	}


def _record_belongs_to_employee_scope(record: dict, scope: dict) -> bool:
	if _is_grand_total_record(record):
		return False
	if scope.get("employee") and record.get("employee") == scope["employee"]:
		return True
	eid = (record.get("employee_id_number") or "").strip()
	if eid and eid in scope["id_variants"]:
		return True
	if eid and eid.upper() in scope["id_variants_upper"]:
		return True
	name_key = _normalize_person_name(record.get("employee_name"))
	if name_key and name_key in scope["name_keys"]:
		return True
	account = (record.get("account_number") or "").strip()
	if account and account in scope.get("account_numbers", set()) and _is_lumpsum_record(record):
		return True
	return False


def _fetch_imported_records_for_scope(
	scope: dict, month_filters: dict, fetch_limit: int | None
) -> list[dict]:
	or_filters: list[list] = []
	if scope.get("employee"):
		or_filters.append(["employee", "=", scope["employee"]])
	for eid in scope.get("id_variants") or []:
		or_filters.append(["employee_id_number", "=", eid])
	if not or_filters:
		or_filters.append(["employee_id_number", "=", ""])

	filters: dict = {"docstatus": 1, **month_filters}
	records = _payslip_get_all(
		"Imported Payroll Record",
		filters=filters,
		or_filters=or_filters,
		fields=PAYSLIP_FIELDS,
		order_by="payroll_month desc, modified desc",
		limit=fetch_limit,
	)
	records = [row for row in records if _record_belongs_to_employee_scope(row, scope)]

	seen = {row.get("name") for row in records}
	accounts = scope.get("account_numbers") or set()
	if accounts and month_filters:
		lumpsum_filters = {
			"docstatus": 1,
			"payroll_sheet_type": "Lumpsum",
			**month_filters,
			"account_number": ["in", list(accounts)],
		}
		for row in _payslip_get_all(
			"Imported Payroll Record",
			filters=lumpsum_filters,
			fields=PAYSLIP_FIELDS,
			order_by="modified desc",
			limit=50,
		):
			if row.get("name") in seen:
				continue
			if _is_grand_total_record(row):
				continue
			seen.add(row.get("name"))
			records.append(row)

	records.sort(
		key=lambda row: (
			getdate(row.get("payroll_month")).isoformat() if row.get("payroll_month") else "",
			str(row.get("modified") or ""),
		),
		reverse=True,
	)
	return records


def _record_sort_key(record: dict) -> tuple:
	modified = record.get("modified")
	return (
		str(modified) if modified else "",
		record.get("name") or "",
	)


def _pick_latest_per_month_and_sheet(records: list[dict]) -> dict[str, dict[str, dict]]:
	by_month: dict[str, dict[str, dict]] = {}
	for record in sorted(records, key=_record_sort_key, reverse=True):
		if _is_grand_total_record(record):
			continue
		month = record.get("payroll_month")
		if not month:
			continue
		month_key = getdate(month).isoformat()
		month_bucket = by_month.setdefault(month_key, {})
		if _is_variable_record(record):
			month_bucket.setdefault("variables", []).append(record)
			continue
		sheet_key = "lumpsum" if _is_lumpsum_record(record) else "staff"
		month_bucket.setdefault(sheet_key, record)
	return by_month


def _merge_staff_and_lumpsum(staff: dict | None, lumpsum: dict | None, month_key: str) -> dict:
	seed = staff or lumpsum or {}
	merged = _new_merged_payslip_row(seed, month_key)
	if staff:
		_apply_staff_record_to_merged(merged, staff)
	if lumpsum:
		_apply_lumpsum_record_to_merged(merged, lumpsum)
	merged["gross_variable"] = 0.0
	merged["variable_labels"] = []
	merged["gross_salary"] = merged["gross_staff"] + merged["gross_lumpsum"]
	merged["net_salary_staff"] = _sheet_net_paid(staff) if staff else 0.0
	merged["net_salary_lumpsum"] = _sheet_net_paid(lumpsum) if lumpsum else 0.0
	_set_payslip_net_from_paid(merged)
	merged["payroll_period_label"] = _payslip_pay_period_label(merged, month_key)
	return merged


def _apply_staff_record_to_merged(target: dict, record: dict) -> None:
	target["basic_salary"] = _flt(record.get("basic_salary"))
	target["transport_allowance"] = _flt(record.get("transport_allowance"))
	target["ict_allowance"] = _flt(record.get("ict_allowance"))
	target["itc_allowance"] = _flt(record.get("itc_allowance"))
	target["gross_staff"] = _flt(record.get("gross_salary"))
	target["paye"] += _flt(record.get("paye"))
	target["pension_employee"] += _flt(record.get("pension_employee"))
	target["maternity_employee"] += _flt(record.get("maternity_employee"))
	target["maternity_employer"] += _flt(record.get("maternity_employer"))
	target["cbhi"] += _flt(record.get("cbhi"))
	_add_other_deductions(target, record)
	_add_net_paid(target, record)
	for field in (
		"employee_name",
		"position",
		"department",
		"company",
		"payroll_excel_upload",
		"account_number",
		"bank",
		"payment_date",
	):
		if record.get(field):
			target[field] = record.get(field)


def _apply_variable_records_to_merged(target: dict, variables: list[dict]) -> None:
	"""Add bonuses and incentives the employee received in this month."""
	if not variables:
		target.setdefault("gross_variable", 0.0)
		target.setdefault("bonus_incentives", 0.0)
		target.setdefault("variable_labels", [])
		return
	labels: list[str] = list(target.get("variable_labels") or [])
	gross_variable = _flt(target.get("gross_variable"))
	bonus_incentives = _flt(target.get("bonus_incentives"))
	variable_net = 0.0
	for record in variables:
		# Basic salary on the payslip stays the staff figure.
		# The variable sheet's Basic salary is the bonus column.
		bonus_incentives += _flt(record.get("basic_salary"))
		target["transport_allowance"] += _flt(record.get("transport_allowance"))
		target["itc_allowance"] += _flt(record.get("itc_allowance"))
		target["ict_allowance"] += _flt(record.get("ict_allowance"))
		gross_variable += _flt(record.get("gross_salary"))
		variable_net += _sheet_net_paid(record)
		target["paye"] += _flt(record.get("paye"))
		target["pension_employee"] += _flt(record.get("pension_employee"))
		target["maternity_employee"] += _flt(record.get("maternity_employee"))
		target["maternity_employer"] += _flt(record.get("maternity_employer"))
		target["cbhi"] += _flt(record.get("cbhi"))
		_add_other_deductions(target, record)
		_add_net_paid(target, record)
		label = (record.get("payroll_period_label") or "Variable").strip()
		if label and label not in labels:
			labels.append(label)
		for field in (
			"employee_name",
			"position",
			"department",
			"company",
			"account_number",
			"bank",
		):
			if record.get(field) and not target.get(field):
				target[field] = record.get(field)
	target["gross_variable"] = gross_variable
	target["bonus_incentives"] = bonus_incentives
	target["variable_labels"] = labels
	target["gross_salary"] = _flt(target.get("gross_staff")) + _flt(target.get("gross_lumpsum")) + gross_variable
	_set_payslip_net_from_paid(target)
	target["net_salary_variable"] = variable_net


def _apply_lumpsum_record_to_merged(target: dict, record: dict) -> None:
	target["gross_lumpsum"] = _flt(record.get("gross_salary"))
	target["paye"] += _flt(record.get("paye"))
	target["pension_employee"] += _flt(record.get("pension_employee"))
	target["maternity_employee"] += _flt(record.get("maternity_employee"))
	target["maternity_employer"] += _flt(record.get("maternity_employer"))
	target["cbhi"] += _flt(record.get("cbhi"))
	_add_other_deductions(target, record)
	_add_net_paid(target, record)
	for field in (
		"employee_name",
		"position",
		"department",
		"company",
		"payroll_excel_upload",
		"account_number",
		"bank",
		"payment_date",
	):
		if record.get(field) and not target.get(field):
			target[field] = record.get(field)


def _payslip_pay_period_label(record: dict, month_key: str | None = None) -> str:
	month = normalize_payroll_month(month_key or record.get("payroll_month"))
	if month:
		return payroll_period_label_for_month(month) or month.strftime("%B %Y")
	raw = (record.get("payroll_period_label") or "").strip()
	if not raw:
		return ""
	# Prefer "September 2026" parsed from Excel titles like "STAFF PAYROLL SEPTEMBER 2026".
	match = re.search(
		r"(January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{4}",
		raw,
		re.I,
	)
	if match:
		return match.group(0).title()
	return raw


def _merge_records_by_payroll_month(records: list[dict], scope: dict | None = None) -> list[dict]:
	scope = scope or {
		"id_variants": set(),
		"id_variants_upper": set(),
		"name_keys": set(),
		"account_numbers": set(),
	}
	filtered = [row for row in records if _record_belongs_to_employee_scope(row, scope)]
	by_month = _pick_latest_per_month_and_sheet(filtered)

	merged_list: list[dict] = []
	for month_key, bucket in by_month.items():
		staff = bucket.get("staff")
		lumpsum = bucket.get("lumpsum")
		variables = bucket.get("variables") or []
		if staff and not lumpsum:
			lumpsum = _match_lumpsum_for_staff(staff, filtered)
		if lumpsum and not staff:
			staff = _match_staff_for_lumpsum(lumpsum, filtered)
		if not staff and not lumpsum and not variables:
			continue
		merged = _merge_staff_and_lumpsum(staff, lumpsum, month_key)
		if not staff and not lumpsum and variables:
			merged = _new_merged_payslip_row(variables[0], month_key)
			merged["gross_variable"] = 0.0
			merged["variable_labels"] = []
		_apply_variable_records_to_merged(merged, variables)
		merged_list.append(merged)

	for row in merged_list:
		row["payroll_period_label"] = _payslip_pay_period_label(row, row.get("payroll_month"))

	merged_list.sort(key=lambda row: row.get("payroll_month") or "", reverse=True)
	return merged_list


def _match_lumpsum_for_staff(staff: dict, records: list[dict]) -> dict | None:
	month_key = getdate(staff.get("payroll_month")).isoformat() if staff.get("payroll_month") else None
	staff_account = (staff.get("account_number") or "").strip()
	for record in sorted(records, key=_record_sort_key, reverse=True):
		if not _is_lumpsum_record(record):
			continue
		if month_key and getdate(record.get("payroll_month")).isoformat() != month_key:
			continue
		if staff.get("employee") and record.get("employee") == staff.get("employee"):
			return record
		if staff_account and (record.get("account_number") or "").strip() == staff_account:
			return record
	return None


def _match_staff_for_lumpsum(lumpsum: dict, records: list[dict]) -> dict | None:
	month_key = getdate(lumpsum.get("payroll_month")).isoformat() if lumpsum.get("payroll_month") else None
	lumpsum_account = (lumpsum.get("account_number") or "").strip()
	for record in sorted(records, key=_record_sort_key, reverse=True):
		if _is_lumpsum_record(record) or _is_variable_record(record):
			continue
		if month_key and getdate(record.get("payroll_month")).isoformat() != month_key:
			continue
		if lumpsum.get("employee") and record.get("employee") == lumpsum.get("employee"):
			return record
		if lumpsum_account and (record.get("account_number") or "").strip() == lumpsum_account:
			return record
	return None


def _new_merged_payslip_row(seed: dict, month_key: str) -> dict:
	return {
		"payroll_month": month_key,
		"payroll_period_label": seed.get("payroll_period_label"),
		"employee_name": seed.get("employee_name"),
		"position": seed.get("position"),
		"department": seed.get("department"),
		"company": seed.get("company"),
		"payroll_excel_upload": seed.get("payroll_excel_upload"),
		"payment_date": seed.get("payment_date"),
		"account_number": seed.get("account_number"),
		"bank": seed.get("bank"),
		"basic_salary": 0.0,
		"transport_allowance": 0.0,
		"bonus_incentives": 0.0,
		"ict_allowance": 0.0,
		"itc_allowance": 0.0,
		"gross_staff": 0.0,
		"gross_lumpsum": 0.0,
		"gross_salary": 0.0,
		"paye": 0.0,
		"pension_employee": 0.0,
		"maternity_employee": 0.0,
		"maternity_employer": 0.0,
		"cbhi": 0.0,
		"muganga_sacco_social_fund": 0.0,
		"rpf": 0.0,
		"brd_loan": 0.0,
		"sport_subscription": 0.0,
		"insurance_sanlam": 0.0,
		"insurance_prime": 0.0,
		"net_salary": 0.0,
		"net_paid": 0.0,
		"take_home": 0.0,
	}


OTHER_DEDUCTION_FIELDS = (
	"muganga_sacco_social_fund",
	"rpf",
	"brd_loan",
	"sport_subscription",
	"insurance_sanlam",
	"insurance_prime",
)


def _other_deductions_amount(record: dict | None) -> float:
	"""Social fund, RPF, loan, sport, and insurance on one payroll sheet."""
	if not record:
		return 0.0
	parts = sum(_flt(record.get(field)) for field in OTHER_DEDUCTION_FIELDS)
	if parts:
		return parts
	return _flt(record.get("other_deductions"))


def _sheet_net_paid(record: dict | None) -> float:
	"""Net paid is take-home after other deductions. Sheets without that column use their net."""
	if not record:
		return 0.0
	paid = _flt(record.get("net_paid"))
	if paid:
		return paid
	return _flt(record.get("net_salary")) or _flt(record.get("take_home"))


def _add_other_deductions(target: dict, record: dict) -> None:
	for field in OTHER_DEDUCTION_FIELDS:
		target[field] = _flt(target.get(field)) + _flt(record.get(field))


def _add_net_paid(target: dict, record: dict) -> None:
	target["net_paid"] = _flt(target.get("net_paid")) + _sheet_net_paid(record)


def _set_payslip_net_from_paid(target: dict) -> None:
	target["net_salary"] = _flt(target.get("net_paid"))
	target["take_home"] = target["net_salary"]


def _payslip_period_hint(
	payroll_month: str | None = None,
	from_month: str | None = None,
	to_month: str | None = None,
	months: int | None = None,
) -> str:
	if payroll_month:
		m = normalize_payroll_month(payroll_month)
		return f" ({m.strftime('%B %Y')})" if m else ""
	if from_month and to_month:
		start = normalize_payroll_month(from_month)
		end = normalize_payroll_month(to_month)
		if start and end:
			return f" ({start.strftime('%B %Y')} – {end.strftime('%B %Y')})"
	if from_month:
		start = normalize_payroll_month(from_month)
		return f" (from {start.strftime('%B %Y')})" if start else ""
	if to_month:
		end = normalize_payroll_month(to_month)
		return f" (to {end.strftime('%B %Y')})" if end else ""
	if months:
		return f" (last {int(months)} months)"
	return ""


def get_last_n_payroll_records(employee_id_number: str, months: int = 3) -> list[dict]:
	return get_payroll_records_for_payslip(employee_id_number, months=months)


def _insert_paragraph_after(paragraph, text: str):
	from docx.text.paragraph import Paragraph
	from docx.oxml import OxmlElement

	new_p = OxmlElement("w:p")
	paragraph._p.addnext(new_p)
	new_para = Paragraph(new_p, paragraph._parent)
	new_para = Paragraph(new_p, paragraph._parent)
	new_para.add_run(text)
	return new_para


def _update_employee_header(doc, record: dict) -> None:
	labels = {
		"EMPLOYEE NAMES:": record.get("employee_name"),
		"POSITION:": record.get("position"),
		"DEPARTMENT:": record.get("department"),
		"ACCOUNT NUMBER:": record.get("account_number"),
		"BENEFICIARY BANK:": record.get("bank"),
	}
	department_para = None
	for paragraph in doc.paragraphs:
		text = paragraph.text or ""
		for prefix, value in labels.items():
			if text.startswith(prefix):
				display = (value or "").strip() or "-"
				paragraph.text = f"{prefix} {display}"
				if prefix == "DEPARTMENT:":
					department_para = paragraph

	variable_labels = [label for label in (record.get("variable_labels") or []) if label]
	if variable_labels and department_para:
		_insert_paragraph_after(
			department_para,
			"",
		)

	if department_para and not any(
		(p.text or "").startswith("ACCOUNT NUMBER:") for p in doc.paragraphs
	):
		acc = (record.get("account_number") or "").strip() or "-"
		bank = (record.get("bank") or "").strip() or "-"
		acc_para = _insert_paragraph_after(department_para, f"ACCOUNT NUMBER: {acc}")
		_insert_paragraph_after(acc_para, f"BENEFICIARY BANK: {bank}")


def _tc_span(tc) -> int:
	from docx.oxml.ns import qn

	tc_pr = tc.find(qn("w:tcPr"))
	if tc_pr is None:
		return 1
	grid_span = tc_pr.find(qn("w:gridSpan"))
	if grid_span is None:
		return 1
	return int(grid_span.get(qn("w:val")) or 1)


def _set_tc_text(tc, text: str) -> None:
	from docx.oxml.ns import qn

	text_nodes = tc.findall(".//" + qn("w:t"))
	if not text_nodes:
		return
	text_nodes[0].text = text
	for node in text_nodes[1:]:
		node.text = ""


def _find_tc_at_grid_col(tr, grid_col: int):
	col = 0
	for tc in tr.tc_lst:
		span = _tc_span(tc)
		if col == grid_col:
			return tc
		col += span
	return None


def _insert_column_after(table, after_grid_col: int, label: str, width: int) -> None:
	"""Insert one column immediately after a grid column. Header text is set on row 0."""
	from docx.oxml.ns import qn

	tbl = table._tbl
	grid = tbl.find(qn("w:tblGrid"))
	cols = list(grid.findall(qn("w:gridCol")))
	if after_grid_col < 0 or after_grid_col >= len(cols):
		return
	new_col = deepcopy(cols[after_grid_col])
	new_col.set(qn("w:w"), str(width))
	cols[after_grid_col].addnext(new_col)

	for row_idx, tr in enumerate(tbl.tr_lst):
		anchor = _find_tc_at_grid_col(tr, after_grid_col)
		if anchor is None:
			continue
		new_tc = deepcopy(anchor)
		for text_node in new_tc.findall(".//" + qn("w:t")):
			text_node.text = ""
		tc_pr = new_tc.find(qn("w:tcPr"))
		if tc_pr is not None:
			width_node = tc_pr.find(qn("w:tcW"))
			if width_node is not None:
				width_node.set(qn("w:w"), str(width))
		if row_idx == 0 and label:
			_set_tc_text(new_tc, label)
		anchor.addnext(new_tc)


def _remove_grid_column(table, grid_col: int) -> None:
	from docx.oxml.ns import qn

	tbl = table._tbl
	grid = tbl.find(qn("w:tblGrid"))
	cols = list(grid.findall(qn("w:gridCol")))
	if grid_col < 0 or grid_col >= len(cols):
		return
	grid.remove(cols[grid_col])
	for tr in tbl.tr_lst:
		col = 0
		for tc in list(tr.tc_lst):
			span = _tc_span(tc)
			if col <= grid_col < col + span:
				if span == 1:
					tr.remove(tc)
				else:
					tc_pr = tc.find(qn("w:tcPr"))
					grid_span = tc_pr.find(qn("w:gridSpan")) if tc_pr is not None else None
					if grid_span is not None:
						grid_span.set(qn("w:val"), str(span - 1))
				break
			col += span


def _set_grid_header(table, grid_col: int, label: str) -> None:
	if not table._tbl.tr_lst:
		return
	tc = _find_tc_at_grid_col(table._tbl.tr_lst[0], grid_col)
	if tc is not None:
		_set_tc_text(tc, label)


def _payslip_has_ict(records: list[dict]) -> bool:
	return any(_flt(record.get("ict_allowance")) for record in records)


def _prepare_payslip_earning_columns(table, include_ict: bool) -> int:
	"""Place Transport and TC Allowance after Basic Salary.

	ICT Allowance is kept only when a payslip row has an ICT amount.
	Bonus and other Incentives follows those allowance columns.
	"""
	_insert_column_after(table, 2, "Transport allowance", TRANSPORT_COLUMN_WIDTH)
	_insert_column_after(table, 3, "TC Allowance", ALLOWANCE_COLUMN_WIDTH)
	if include_ict:
		bonus_after = 5
		lumpsum_col = 7
	else:
		_remove_grid_column(table, 5)
		bonus_after = 4
		lumpsum_col = 6
	_insert_column_after(table, bonus_after, "Bonus and other Incentives", BONUS_COLUMN_WIDTH)
	_set_grid_header(table, lumpsum_col, "Gross Lumpsum")
	_fit_table_width(table)
	return PAYSLIP_TABLE_COLS_WITH_ICT if include_ict else PAYSLIP_TABLE_COLS_WITHOUT_ICT


def _fit_table_width(table, target: int = 14845) -> None:
	from docx.oxml.ns import qn

	tbl = table._tbl
	grid_cols = list(tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol")))
	current = sum(int(col.get(qn("w:w")) or 0) for col in grid_cols)
	if not current or current == target:
		return
	scaled = [max(80, int(int(col.get(qn("w:w")) or 0) * target / current)) for col in grid_cols]
	scaled[-1] += target - sum(scaled)
	for col, width in zip(grid_cols, scaled):
		col.set(qn("w:w"), str(max(width, 80)))
	for tr in tbl.tr_lst:
		for tc in tr.tc_lst:
			tc_pr = tc.find(qn("w:tcPr"))
			if tc_pr is None:
				continue
			width = tc_pr.find(qn("w:tcW"))
			if width is None or not width.get(qn("w:w")):
				continue
			width.set(qn("w:w"), str(max(80, int(int(width.get(qn("w:w"))) * target / current))))


def _ensure_table_columns(table, target_cols: int) -> None:
	for row in table.rows:
		while len(row.cells) > target_cols:
			row._tr.remove(row.cells[-1]._tc)
		while len(row.cells) < target_cols:
			row._tr.append(deepcopy(row.cells[-1]._tc))


def _trim_table_data_rows(table, first_data_row: int, display_rows: int) -> None:
	target_rows = first_data_row + display_rows
	while len(table.rows) > target_rows:
		table._tbl.remove(table.rows[-1]._tr)


def _apply_one_page_layout(doc, num_data_rows: int) -> None:
	from docx.shared import Inches, Pt

	body_size = Pt(8 if num_data_rows <= 1 else 7)
	for section in doc.sections:
		section.top_margin = Inches(0.35)
		section.bottom_margin = Inches(0.35)
		section.left_margin = Inches(0.4)
		section.right_margin = Inches(0.4)

	for paragraph in doc.paragraphs:
		pf = paragraph.paragraph_format
		pf.space_before = Pt(0)
		pf.space_after = Pt(1)
		for run in paragraph.runs:
			run.font.size = body_size

	for table in doc.tables:
		for row in table.rows:
			row.height = None
			for cell in row.cells:
				for paragraph in cell.paragraphs:
					pf = paragraph.paragraph_format
					pf.space_before = Pt(0)
					pf.space_after = Pt(0)
					for run in paragraph.runs:
						run.font.size = body_size


def build_payslip_docx(
	employee_id_number: str,
	months: int = 3,
	from_month: str | None = None,
	to_month: str | None = None,
	payroll_month: str | None = None,
	signoff: dict | None = None,
) -> bytes:
	try:
		from docx import Document
	except ImportError as e:
		frappe.throw(_("python-docx is required. Install: pip install python-docx"))

	records = get_payroll_records_for_payslip(
		employee_id_number,
		months=months,
		from_month=from_month,
		to_month=to_month,
		payroll_month=payroll_month,
	)
	if not records:
		employee_id_number = normalize_employee_id_number(employee_id_number)
		period_hint = _payslip_period_hint(
			payroll_month=payroll_month,
			from_month=from_month,
			to_month=to_month,
			months=months,
		)
		latest_rows = _payslip_get_all(
			"Imported Payroll Record",
			filters={"employee_id_number": employee_id_number, "docstatus": 1},
			fields=["payroll_month"],
			order_by="payroll_month desc",
			limit=1,
		)
		latest = latest_rows[0].payroll_month if latest_rows else None
		msg = _("No submitted payroll records found for employee ID {0}{1}.").format(
			employee_id_number,
			period_hint,
		)
		if latest:
			msg += " " + _("Latest imported payroll month for this ID is {0}.").format(
				getdate(latest).strftime("%B %Y")
			)
		frappe.throw(msg)

	if payroll_month or (from_month and to_month) or from_month or to_month:
		display_rows = len(records)
	else:
		display_rows = min(int(months or 3), MAX_PAYSLIP_MONTHS)
	display_rows = min(display_rows, MAX_PAYSLIP_ROWS_ONE_PAGE)
	records = records[:display_rows]

	template_path = get_template_path()
	if not os.path.exists(template_path):
		frappe.throw(_("Payslip template not found at {0}").format(template_path))

	doc = Document(template_path)
	latest = records[0]

	_inject_payslip_logo(doc, latest)
	_update_employee_header(doc, latest)

	if not doc.tables:
		frappe.throw(_("Payslip template must contain a payroll table."))

	table = doc.tables[0]
	first_data_row = 2
	include_ict = _payslip_has_ict(records)
	table_cols = _prepare_payslip_earning_columns(table, include_ict)

	while len(table.rows) < first_data_row + display_rows:
		table.add_row()
	_ensure_table_columns(table, table_cols)

	for row_idx in range(first_data_row, len(table.rows)):
		for cell in table.rows[row_idx].cells:
			cell.text = "-"

	for idx, record in enumerate(records):
		row = table.rows[first_data_row + idx]
		values = _record_to_table_row(record, include_ict=include_ict)
		for col_idx, value in enumerate(values):
			if col_idx < len(row.cells):
				row.cells[col_idx].text = value

	_trim_table_data_rows(table, first_data_row, display_rows)
	_apply_one_page_layout(doc, display_rows)
	_inject_employee_photo(doc, latest, employee_id_number)
	_append_payslip_signoff(doc, signoff)

	buffer = io.BytesIO()
	doc.save(buffer)
	buffer.seek(0)
	return buffer.getvalue()


def _record_to_table_row(record: dict, include_ict: bool = False) -> list[str]:
	pay_period = _payslip_pay_period_label(record)
	payment_date = record.get("payment_date")
	if payment_date:
		payment_date = getdate(payment_date).strftime("%d-%m-%Y")
	elif record.get("payroll_month"):
		payment_date = _default_payment_date(record.get("payroll_month"))
	else:
		payment_date = "-"

	gross_lumpsum = _flt(record.get("gross_lumpsum"))
	maternity = _flt(record.get("maternity_employee")) + _flt(record.get("maternity_employer"))
	other_deductions = _other_deductions_amount(record)
	gross_salary = _flt(record.get("gross_salary"))
	net = _sheet_net_paid(record)
	total_deductions = gross_salary - net if gross_salary and net else 0.0

	values = [
		pay_period,
		payment_date,
		_fmt(record.get("basic_salary")),
		_fmt(record.get("transport_allowance")),
		_fmt(record.get("itc_allowance")),
	]
	if include_ict:
		values.append(_fmt(record.get("ict_allowance")))
	values.extend([
		_fmt(record.get("bonus_incentives")),
		_fmt(gross_lumpsum) if gross_lumpsum else "-",
		_fmt(gross_salary),
		_fmt(record.get("paye")),
		_fmt(record.get("pension_employee")),
		_fmt(maternity / 2) if maternity else "-",
		_fmt(maternity / 2) if maternity else "-",
		_fmt(record.get("cbhi")),
		_fmt(other_deductions) if other_deductions else "-",
		_fmt(total_deductions),
		_fmt(net),
	])
	return values


def _txt(val) -> str:
	text = (val or "").strip()
	return text if text else "-"


def _default_payment_date(payroll_month) -> str:
	d = getdate(payroll_month)
	# Typical pay date: 25th of same month
	pay_day = date(d.year, d.month, min(25, _days_in_month(d.year, d.month)))
	return pay_day.strftime("%d-%m-%Y")


def _days_in_month(year: int, month: int) -> int:
	if month == 12:
		next_month = date(year + 1, 1, 1)
	else:
		next_month = date(year, month + 1, 1)
	return (next_month - date(year, month, 1)).days


def _fmt(val) -> str:
	val = _flt(val)
	if not val:
		return "-"
	return f"{val:,.0f}".replace(",", ",")


def _flt(val):
	if val is None or val == "":
		return 0.0
	try:
		return float(val)
	except (TypeError, ValueError):
		return 0.0
