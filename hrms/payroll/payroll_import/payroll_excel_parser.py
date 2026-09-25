# Copyright (c) 2026, Frappe Technologies Pvt. Ltd. and contributors
# License: GNU General Public License v3. See license.txt

from __future__ import annotations

import re
from datetime import date
from typing import Any

import frappe
from frappe import _
from frappe.utils import getdate


MONTH_MAP = {
	"january": 1,
	"february": 2,
	"march": 3,
	"april": 4,
	"may": 5,
	"june": 6,
	"july": 7,
	"august": 8,
	"september": 9,
	"october": 10,
	"november": 11,
	"december": 12,
}


def parse_payroll_excel(file_path: str) -> dict[str, Any]:
	try:
		from openpyxl import load_workbook
	except ImportError as e:
		frappe.throw(_("openpyxl is required. Install it in the bench env: pip install openpyxl"))

	wb = load_workbook(file_path, read_only=True, data_only=True)
	sheet_name = wb.sheetnames[0]
	ws = wb[sheet_name]

	rows = list(ws.iter_rows(values_only=True))
	wb.close()

	title_row, header_row_idx, sub_header_rows, data_start = _find_header_and_data_start(rows)
	payroll_period_label, payroll_month = _parse_period_from_title(title_row)
	if not payroll_month:
		payroll_period_label, payroll_month = _parse_period_from_rows(rows, header_row_idx)
	headers = _normalize_headers(rows[header_row_idx])
	column_map = _build_column_map(headers, sub_header_rows)
	by_key = _headers_to_by_key(headers, sub_header_rows)

	while data_start < len(rows) and not _is_data_row(rows[data_start], column_map):
		data_start += 1

	records = []
	for row in rows[data_start :]:
		if not row or not _is_data_row(row, column_map):
			continue
		records.append(_row_to_record(row, column_map, payroll_period_label, payroll_month))

	sheet_footer = parse_sheet_footer_from_rows(rows)
	payroll_sheet_type = detect_payroll_sheet_type(rows)

	return {
		"sheet_name": sheet_name,
		"sheet_title": title_row,
		"payroll_period_label": payroll_period_label,
		"payroll_month": payroll_month,
		"payroll_sheet_type": payroll_sheet_type,
		"column_map": column_map,
		"unmapped_headers": unmapped_excel_headers(by_key, column_map),
		"records": records,
		"sheet_footer": sheet_footer,
	}


def parse_sheet_footer_from_rows(rows: list[tuple]) -> dict | None:
	"""Read signature block and date line from the bottom of the payroll workbook."""
	date_line = None
	signature_row_idx = None

	for row in rows:
		for cell in row:
			if not isinstance(cell, str):
				continue
			text = cell.strip()
			if text.lower().startswith("done at"):
				date_line = text
				break

	for idx, row in enumerate(rows):
		for cell in row:
			if isinstance(cell, str) and "prepared by" in cell.lower():
				signature_row_idx = idx
				break
		if signature_row_idx is not None:
			break

	if signature_row_idx is None:
		return {"date_line": date_line} if date_line else None

	headings_row = rows[signature_row_idx]
	names_row = rows[signature_row_idx + 1] if signature_row_idx + 1 < len(rows) else ()
	titles_row = rows[signature_row_idx + 2] if signature_row_idx + 2 < len(rows) else ()

	signatures = []
	for col_idx, cell in enumerate(headings_row):
		heading = str(cell or "").strip()
		if not heading:
			continue
		lower = heading.lower()
		if not any(k in lower for k in ("prepared by", "verified by", "approved by")):
			continue
		if not heading.endswith(":"):
			heading = f"{heading}:"
		name = ""
		title = ""
		if col_idx < len(names_row) and names_row[col_idx] is not None:
			name = str(names_row[col_idx]).strip()
		if col_idx < len(titles_row) and titles_row[col_idx] is not None:
			title = str(titles_row[col_idx]).strip()
		signatures.append({"heading": heading, "name": name, "title": title})

	if not date_line and not signatures:
		return None

	return {"date_line": date_line, "signatures": signatures}


def detect_payroll_sheet_type(rows: list[tuple]) -> str:
	for row in rows:
		if not row:
			continue
		for cell in row:
			if isinstance(cell, str) and "LUMPSUM" in cell.upper():
				return "Lumpsum"
	return "Staff Payroll"


def _find_header_and_data_start(rows: list[tuple]) -> tuple[str | None, int, list[tuple], int]:
	title = None
	header_idx = None
	for idx, row in enumerate(rows):
		if not row:
			continue
		for cell in row:
			if isinstance(cell, str):
				upper = cell.upper()
				if "STAFF PAYROLL" in upper:
					title = cell.strip()
				elif "LUMPSUM" in upper and "GROSS" not in upper:
					candidate = cell.strip()
					if candidate and (not title or len(candidate) > len(title)):
						title = candidate
		first_cell = str(row[0]).strip().upper() if row[0] is not None else ""
		if first_cell in ("S/N", "S/N ", "SN", "NAMES", "NAME", "#"):
			header_idx = idx
			break
		joined = " ".join(str(c or "").strip().upper() for c in row[:6])
		if "NAMES" in joined and ("GROSS" in joined or "LUMPSUM" in joined or "DEPARTMENT" in joined):
			header_idx = idx
			break
	if header_idx is None:
		frappe.throw(_("Could not find payroll header row (S/N or NAMES) in the Excel file."))

	sub_header_rows: list[tuple] = []
	data_start = header_idx + 1
	while data_start < len(rows):
		row = rows[data_start]
		if not row:
			data_start += 1
			continue
		first = str(row[0] or "").strip().upper()
		if first in ("", "S/N", "SN", "NAMES", "NAME") or "PENSION" in first or "RSSB" in first:
			sub_header_rows.append(row)
			data_start += 1
			continue
		break

	return title, header_idx, sub_header_rows, data_start


def _parse_period_from_rows(rows: list[tuple], header_row_idx: int) -> tuple[str, str | None]:
	for row in rows[: header_row_idx + 1]:
		if not row:
			continue
		for cell in row:
			if not isinstance(cell, str):
				continue
			text = cell.strip()
			label, month = _parse_period_from_title(text)
			if month:
				return label or text, month
	return "", None


def _parse_period_from_title(title: str | None) -> tuple[str, str | None]:
	if not title:
		return "", None
	match = re.search(
		r"(January|February|March|April|May|June|July|August|September|October|November|December)\s+(\d{4})",
		title,
		re.I,
	)
	if not match:
		return title, None
	month_name, year = match.group(1), int(match.group(2))
	month_num = MONTH_MAP[month_name.lower()]
	payroll_month = date(year, month_num, 1).isoformat()
	label = title.strip()
	return label, payroll_month


def _normalize_headers(header_row: tuple) -> list[str]:
	return [str(c).strip() if c is not None else "" for c in header_row]


def _fill_merged_headers(headers: list[str]) -> list[str]:
	filled = []
	last = ""
	for h in headers:
		h = (h or "").strip()
		if h:
			last = h
		filled.append(last)
	return filled


def _register_column_keys(mapping: dict[str, int], idx: int, *labels: str) -> None:
	for label in labels:
		if not label:
			continue
		key = _header_key(label)
		if key:
			mapping.setdefault(key, idx)


def _headers_to_by_key(headers: list[str], sub_header_rows: list[tuple]) -> dict[str, int]:
	filled = _fill_merged_headers(headers)
	n = len(filled)
	for sub_row in sub_header_rows:
		n = max(n, len(sub_row))

	by_key: dict[str, int] = {}
	for idx in range(n):
		parent = filled[idx] if idx < len(filled) else ""
		sub_labels: list[str] = []
		for sub_row in sub_header_rows:
			if idx < len(sub_row) and sub_row[idx] is not None and str(sub_row[idx]).strip():
				sub_labels.append(str(sub_row[idx]).strip())
		sub_val = sub_labels[-1] if sub_labels else ""
		combined = " ".join([p for p in [parent, *sub_labels] if p])
		if combined:
			_register_column_keys(by_key, idx, combined, sub_val, parent)
		elif sub_val:
			_register_column_keys(by_key, idx, sub_val)
		elif parent:
			_register_column_keys(by_key, idx, parent)
	return by_key


def _build_column_map(headers: list[str], sub_header_rows: list[tuple]) -> dict[str, int]:
	by_key = _headers_to_by_key(headers, sub_header_rows)

	aliases = {
		"names": "employee_name",
		"name": "employee_name",
		"id_number": "employee_id_number",
		"id_no": "employee_id_number",
		"basic_salary": "basic_salary",
		"basic": "basic_salary",
		"gross_salary": "gross_salary",
		"gross_lumpsum": "gross_salary",
		"gross": "gross_salary",
		"department": "department",
		"paye": "paye",
		"paye_pay_as_you_earn": "paye",
		"net_salary": "net_salary",
		"net_paid": "net_paid",
		"take_home": "take_home",
		"position": "position",
		"date_of_joining": "date_of_joining",
		"transport_allowance": "transport_allowance",
		"transport_allowance_cro": "transport_allowance",
		"itc_allowance": "itc_allowance",
		"ict_allowance": "ict_allowance",
		"net_before_cbhi": "net_before_cbhi",
		"rssb_cbhi_scheme_0_5": "cbhi",
		"rssb_cbhi": "cbhi",
		"pension_employee_6": "pension_employee",
		"pension_employee_6_": "pension_employee",
		"rssb_contributions_pension_scheme_oh_pension_employee_6": "pension_employee",
		"pension_employee": "pension_employee",
		"pension_employer_6": "pension_employer",
		"rssb_contributions_pension_scheme_oh_pension_employer_6": "pension_employer",
		"pension_employer": "pension_employer",
		"oh_employer_2": "oh_employer",
		"rssb_contributions_pension_scheme_oh_oh_employer_2": "oh_employer",
		"oh_employer": "oh_employer",
		"maternity_employee_0_3": "maternity_employee",
		"employee_0_3": "maternity_employee",
		"employee_0_3_": "maternity_employee",
		"maternity_leave_benefits_employee_0_3": "maternity_employee",
		"maternity_employer_0_3": "maternity_employer",
		"employer_0_3": "maternity_employer",
		"maternity_leave_benefits_employer_0_3": "maternity_employer",
		"muganga_sacco_social_fund": "muganga_sacco_social_fund",
		"rpf": "rpf",
		"brd_scholar_loan_repayment": "brd_loan",
		"sport_subscription_reimbursement": "sport_subscription",
		"sanlam": "insurance_sanlam",
		"prime": "insurance_prime",
		"compulsory_saving": "compulsory_saving",
		"compulsory_saving_account": "compulsory_saving_account",
		"account_number": "account_number",
		"beneficiary_bank": "bank",
	}
	field_map: dict[str, int] = {}
	for alias, field in aliases.items():
		if alias in by_key:
			field_map.setdefault(field, by_key[alias])

	_apply_fuzzy_field_map(by_key, field_map)
	_apply_importable_docfield_map(by_key, field_map)
	return field_map


def _importable_payroll_fieldnames() -> frozenset[str]:
	meta = frappe.get_meta("Imported Payroll Record")
	skip = frozenset(
		{
			"name",
			"owner",
			"creation",
			"modified",
			"modified_by",
			"docstatus",
			"idx",
			"amended_from",
			"employee",
			"payroll_excel_upload",
			"company",
			"currency",
		}
	)
	types = frozenset({"Data", "Date", "Float", "Currency", "Int", "Small Text"})
	return frozenset(
		df.fieldname
		for df in meta.fields
		if df.fieldname not in skip and df.fieldtype in types and not df.read_only
	)


def _apply_importable_docfield_map(by_key: dict[str, int], field_map: dict[str, int]) -> None:
	"""Map Excel headers that match Imported Payroll Record field names (e.g. department)."""
	for fieldname in _importable_payroll_fieldnames():
		if fieldname in field_map:
			continue
		if fieldname in by_key:
			field_map[fieldname] = by_key[fieldname]


def unmapped_excel_headers(by_key: dict[str, int], field_map: dict[str, int]) -> list[str]:
	used = set(field_map.values())
	ignore = frozenset({"sn", "s_n", "s_no"})
	labels: list[str] = []
	for key, idx in sorted(by_key.items(), key=lambda item: item[1]):
		if idx in used or key in ignore:
			continue
		labels.append(key.replace("_", " ").title())
	return labels


def _apply_fuzzy_field_map(by_key: dict[str, int], field_map: dict[str, int]) -> None:
	rules: list[tuple[str, Any]] = [
		(
			"compulsory_saving_account",
			lambda k: "compulsory" in k and "account" in k,
		),
		(
			"compulsory_saving",
			lambda k: "compulsory" in k and "saving" in k and "account" not in k,
		),
		("transport_allowance", lambda k: "transport" in k and "allowance" in k),
		("transport_allowance", lambda k: k in ("cro",) or k.endswith("_cro")),
		("itc_allowance", lambda k: "itc" in k and "allowance" in k),
		("ict_allowance", lambda k: "ict" in k and "allowance" in k),
		("paye", lambda k: "paye" in k),
		("pension_employee", lambda k: "pension" in k and "employee" in k),
		("pension_employer", lambda k: "pension" in k and "employer" in k and "oh" not in k),
		("oh_employer", lambda k: ("oh" in k and "employer" in k) or k.startswith("oh_")),
		("maternity_employee", lambda k: "maternity" in k and "employee" in k),
		("maternity_employer", lambda k: "maternity" in k and "employer" in k),
		("net_before_cbhi", lambda k: "net" in k and "before" in k and "cbhi" in k),
		("cbhi", lambda k: "cbhi" in k and "before" not in k),
		("basic_salary", lambda k: k == "basic" or ("basic" in k and "salary" in k)),
		("gross_salary", lambda k: k == "gross" or ("gross" in k and "salary" in k)),
		("gross_salary", lambda k: "gross" in k and "lumpsum" in k),
		("department", lambda k: k == "department" or k == "dept"),
		("date_of_joining", lambda k: "joining" in k and "date" in k),
		("employee_id_number", lambda k: k in ("id_number", "id_no") or ("id" in k and "number" in k)),
		("employee_name", lambda k: k in ("names", "name") and "account" not in k and "bank" not in k),
		("position", lambda k: k == "position"),
		("account_number", lambda k: "account" in k and "number" in k),
		("bank", lambda k: "beneficiary" in k and "bank" in k),
	]
	for field, predicate in rules:
		if field in field_map:
			continue
		for key, idx in sorted(by_key.items(), key=lambda item: (-len(item[0]), item[0])):
			if predicate(key):
				field_map[field] = idx
				break


def _header_key(text: str) -> str:
	text = text.strip().lower()
	if not text:
		return ""
	text = re.sub(r"[^a-z0-9]+", "_", text)
	return text.strip("_")


def _is_footer_or_signature_row(row: tuple) -> bool:
	text = " ".join(str(cell or "").strip().lower() for cell in row[:8])
	if not text:
		return True
	markers = (
		"done at",
		"prepared by",
		"verified by",
		"approved by",
		"grand total",
		"signature",
	)
	return any(marker in text for marker in markers)


def _is_data_row(row: tuple, column_map: dict[str, int] | None = None) -> bool:
	if not row:
		return False
	if _is_footer_or_signature_row(row):
		return False
	for cell in row:
		if isinstance(cell, str) and cell.strip().upper().startswith("GRAND TOTAL"):
			return False

	first = row[0]
	if first is None:
		return _row_has_payroll_amounts(row, column_map)
	if isinstance(first, (int, float)) and not isinstance(first, bool):
		return True
	if isinstance(first, str):
		val = first.strip()
		if not val:
			return _row_has_payroll_amounts(row, column_map)
		if val.isdigit():
			return True
		if val.upper().startswith("GRAND TOTAL"):
			return False
		if _row_has_payroll_amounts(row, column_map):
			return True
	return _row_has_payroll_amounts(row, column_map)


def _row_has_payroll_amounts(row: tuple, column_map: dict[str, int] | None) -> bool:
	if not column_map:
		return False
	for field in ("gross_salary", "net_salary", "paye", "take_home"):
		idx = column_map.get(field)
		if idx is None or idx >= len(row):
			continue
		if _flt(row[idx]):
			return True
	has_identity = False
	for field in ("employee_id_number", "employee_name"):
		idx = column_map.get(field)
		if idx is None or idx >= len(row):
			continue
		val = row[idx]
		if val is None or val == "":
			continue
		text = str(val).strip()
		if text and not text.upper().startswith("GRAND TOTAL"):
			has_identity = True
	return has_identity and any(_flt(row[idx]) for idx in (column_map.get(f) for f in ("gross_salary", "net_salary", "paye", "take_home")) if idx is not None and idx < len(row))


def _cell(row: tuple, column_map: dict[str, int], field: str, default=None):
	idx = column_map.get(field)
	if idx is None or idx >= len(row):
		return default
	return row[idx]


def _row_to_record(
	row: tuple, column_map: dict[str, int], payroll_period_label: str, payroll_month: str | None
) -> dict[str, Any]:
	doj = _cell(row, column_map, "date_of_joining")
	if doj:
		if isinstance(doj, str) and re.match(r"^\d{1,2}-\d{1,2}-\d{4}$", doj.strip()):
			parts = doj.strip().split("-")
			doj = date(int(parts[2]), int(parts[1]), int(parts[0])).isoformat()
		else:
			doj = getdate(doj).isoformat()

	record = {
		"employee_id_number": _cell(row, column_map, "employee_id_number"),
		"employee_name": _cell(row, column_map, "employee_name"),
		"date_of_joining": doj,
		"position": _cell(row, column_map, "position"),
		"department": _cell(row, column_map, "department"),
		"payroll_period_label": payroll_period_label,
		"payroll_month": payroll_month,
		"basic_salary": _flt(_cell(row, column_map, "basic_salary")),
		"transport_allowance": _flt(_cell(row, column_map, "transport_allowance")),
		"itc_allowance": _flt(_cell(row, column_map, "itc_allowance")),
		"ict_allowance": _flt(_cell(row, column_map, "ict_allowance")),
		"gross_salary": _flt(_cell(row, column_map, "gross_salary")),
		"paye": _flt(_cell(row, column_map, "paye")),
		"pension_employee": _flt(_cell(row, column_map, "pension_employee")),
		"pension_employer": _flt(_cell(row, column_map, "pension_employer")),
		"oh_employer": _flt(_cell(row, column_map, "oh_employer")),
		"maternity_employee": _flt(_cell(row, column_map, "maternity_employee")),
		"maternity_employer": _flt(_cell(row, column_map, "maternity_employer")),
		"net_before_cbhi": _flt(_cell(row, column_map, "net_before_cbhi")),
		"cbhi": _flt(_cell(row, column_map, "cbhi")),
		"net_salary": _flt(_cell(row, column_map, "net_salary")),
		"muganga_sacco_social_fund": _flt(_cell(row, column_map, "muganga_sacco_social_fund")),
		"rpf": _flt(_cell(row, column_map, "rpf")),
		"brd_loan": _flt(_cell(row, column_map, "brd_loan")),
		"sport_subscription": _flt(_cell(row, column_map, "sport_subscription")),
		"insurance_sanlam": _flt(_cell(row, column_map, "insurance_sanlam")),
		"insurance_prime": _flt(_cell(row, column_map, "insurance_prime")),
		"net_paid": _flt(_cell(row, column_map, "net_paid")),
		"compulsory_saving": _flt(_cell(row, column_map, "compulsory_saving")),
		"compulsory_saving_account": _cell(row, column_map, "compulsory_saving_account"),
		"take_home": _flt(_cell(row, column_map, "take_home")),
		"account_number": _cell(row, column_map, "account_number"),
		"bank": _cell(row, column_map, "bank"),
	}
	record["other_deductions"] = _flt(record["muganga_sacco_social_fund"]) + _flt(record["rpf"]) + _flt(
		record["brd_loan"]
	) + _flt(record["sport_subscription"]) + _flt(record["insurance_sanlam"]) + _flt(record["insurance_prime"])
	record["total_deductions"] = _flt(record["gross_salary"]) - _flt(record["net_salary"])
	if not record["take_home"] and record["net_salary"]:
		record["take_home"] = record["net_salary"]

	for field in _importable_payroll_fieldnames():
		if field in record:
			continue
		if field not in column_map:
			continue
		val = _cell(row, column_map, field)
		if val is None or val == "":
			continue
		fieldtype = frappe.get_meta("Imported Payroll Record").get_field(field).fieldtype
		if fieldtype == "Date":
			val = getdate(val).isoformat()
		elif fieldtype in ("Float", "Currency", "Int"):
			val = _flt(val)
		else:
			val = str(val).strip()
		record[field] = val

	if not record["employee_id_number"] and record.get("employee_name"):
		record["employee_id_number"] = str(record["employee_name"]).strip()
	if not record["employee_id_number"] and not record["employee_name"]:
		frappe.throw(_("Row is missing employee ID and name."))
	return record


def _flt(val):
	if val is None or val == "":
		return 0.0
	try:
		return float(val)
	except (TypeError, ValueError):
		return 0.0
