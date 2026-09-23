# Copyright (c) 2026, Frappe Technologies Pvt. Ltd. and contributors
# License: GNU General Public License v3. See license.txt

from __future__ import annotations

import io
import frappe
from frappe import _
from frappe.utils import add_months, getdate

from hrms.payroll.payroll_import.payroll_sheet_layout import (
	ensure_payroll_sheet_layout,
	get_company_logo_path,
	get_payroll_excel_upload_for_month,
	get_payroll_sheet_footer,
)

# (field, column label, is_numeric)
DATA_COLUMNS: list[tuple[str | None, str, bool]] = [
	("employee_id_number", "ID NUMBER", False),
	("employee_name", "NAMES", False),
	("date_of_joining", "Date of Joining", False),
	("position", "POSITION", False),
	("basic_salary", "Basic salary", True),
	("transport_allowance", "Transport allowance CRO", True),
	("itc_allowance", "ITC allowance", True),
	("ict_allowance", "ICT allowance", True),
	("gross_salary", "Gross Salary", True),
	("paye", "PAYE(PAY AS YOU EARN)", True),
	("pension_employee", "Pension Employee 6%", True),
	("pension_employer", "Pension Employer 6%", True),
	("oh_employer", "OH Employer 2%", True),
	("maternity_employee", "Employee 0,3%", True),
	("maternity_employer", "Employer 0,3%", True),
	("net_before_cbhi", "Net before CBHI", True),
	("cbhi", "RSSB CBHI Scheme 0.5%", True),
	("net_salary", "Net Salary", True),
	("net_salary", "Net Salary", True),
	("muganga_sacco_social_fund", "MUGANGA SACCO SOCIAL FUND", True),
	("rpf", "RPF", True),
	("brd_loan", "BRD Scholar loan Repayment", True),
	("sport_subscription", "Sport Subscription reimbursement", True),
	("insurance_sanlam", "Sanlam", True),
	("insurance_prime", "Prime", True),
	("net_paid", "Net paid", True),
	("net_paid", "Net paid", True),
	("compulsory_saving", "Compulsory Saving", True),
	("compulsory_saving_account", "Compulsory saving Account", False),
	("take_home", "Take home", True),
	("account_number", "Account Number", False),
	("bank", "Beneficiary Bank", False),
]

YELLOW_HEADERS = {"Net Salary", "Net paid"}
GREEN_HEADERS = {"Take home", "Compulsory Saving", "Compulsory saving Account"}
RSSB_LABELS = {
	"Pension Employee 6%",
	"Pension Employer 6%",
	"OH Employer 2%",
	"Employee 0,3%",
	"Employer 0,3%",
}
INSURANCE_LABELS = {"Sanlam", "Prime"}


def build_payroll_month_excel(payroll_month: str, company: str | None = None) -> bytes:
	try:
		from openpyxl import Workbook
		from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
		from openpyxl.utils import get_column_letter
	except ImportError:
		frappe.throw(_("openpyxl is required. Install: pip install openpyxl"))

	month = getdate(payroll_month)
	records = _fetch_records(month, company)
	if not records:
		frappe.throw(_("No submitted payroll records for {0}.").format(month.strftime("%B %Y")))

	period_label = records[0].get("payroll_period_label") or month.strftime("%B %Y")
	prev_month = getdate(add_months(month, -1))
	prev_records = _fetch_records(prev_month, company)

	wb = Workbook()
	ws = wb.active
	ws.title = "STAFF PAYROLL"

	last_col = 1 + len(DATA_COLUMNS)
	thin = Side(style="thin", color="000000")
	border = Border(left=thin, right=thin, top=thin, bottom=thin)
	header_fill = PatternFill("solid", fgColor="D9D9D9")
	sub_fill = PatternFill("solid", fgColor="BFBFBF")
	yellow_fill = PatternFill("solid", fgColor="FFF2CC")
	green_fill = PatternFill("solid", fgColor="C6E0B4")
	title_font = Font(bold=True, size=14)
	header_font = Font(bold=True, size=9)

	payroll_excel_upload = get_payroll_excel_upload_for_month(company, month)
	ensure_payroll_sheet_layout(payroll_excel_upload)
	logo_path = get_company_logo_path(company, payroll_excel_upload)
	if logo_path:
		try:
			from openpyxl.drawing.image import Image as XLImage

			img = XLImage(logo_path)
			img.height = 90
			img.width = 200
			ws.add_image(img, "A1")
			for r in range(1, 7):
				ws.row_dimensions[r].height = 18
		except Exception:
			frappe.log_error(title="Payroll Excel Logo")

	title_row = 7
	h1, h2, h3 = 8, 9, 10
	data_start = 11

	ws.merge_cells(start_row=title_row, start_column=2, end_row=title_row, end_column=last_col)
	title_cell = ws.cell(row=title_row, column=2, value=f"STAFF PAYROLL {period_label.upper()}")
	title_cell.font = title_font
	title_cell.alignment = Alignment(horizontal="center", vertical="center")

	rssb_cols = []
	ins_cols = []
	for idx, (_f, label, _n) in enumerate(DATA_COLUMNS, start=2):
		if label in RSSB_LABELS:
			rssb_cols.append(idx)
		if label in INSURANCE_LABELS:
			ins_cols.append(idx)

	rssb_start, rssb_end = rssb_cols[0], rssb_cols[-1]
	pension_cols = rssb_cols[:3]
	maternity_cols = rssb_cols[3:]
	ins_start, ins_end = ins_cols[0], ins_cols[-1]

	def style_header(cell, fill=None):
		cell.font = header_font
		cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
		cell.border = border
		cell.fill = fill or header_fill

	# S/N spans 3 header rows
	ws.cell(row=h1, column=1, value="S/N")
	ws.merge_cells(start_row=h1, start_column=1, end_row=h3, end_column=1)
	style_header(ws.cell(row=h1, column=1))

	for col in range(2, last_col + 1):
		label = DATA_COLUMNS[col - 2][1]
		in_rssb = col in rssb_cols
		in_ins = col in ins_cols

		if in_rssb:
			if col == rssb_start:
				ws.merge_cells(start_row=h1, start_column=rssb_start, end_row=h1, end_column=rssb_end)
				c = ws.cell(row=h1, column=rssb_start, value="RSSB CONTRIBUTIONS")
				style_header(c)
			if col == pension_cols[0]:
				ws.merge_cells(start_row=h2, start_column=pension_cols[0], end_row=h2, end_column=pension_cols[-1])
				c = ws.cell(row=h2, column=pension_cols[0], value="PENSION SCHEME + OH")
				style_header(c, sub_fill)
			if col == maternity_cols[0]:
				ws.merge_cells(start_row=h2, start_column=maternity_cols[0], end_row=h2, end_column=maternity_cols[-1])
				c = ws.cell(row=h2, column=maternity_cols[0], value="MATERNITY LEAVE BENEFITS")
				style_header(c, sub_fill)
			c = ws.cell(row=h3, column=col, value=label)
			fill = sub_fill if label.startswith(("Employee 0", "Employer 0", "OH Employer")) else header_fill
			style_header(c, fill)
			continue

		if in_ins:
			if col == ins_start:
				ws.merge_cells(start_row=h1, start_column=ins_start, end_row=h2, end_column=ins_end)
				c = ws.cell(row=h1, column=ins_start, value="INSURANCE DEDUCTIONS")
				style_header(c)
			c = ws.cell(row=h3, column=col, value=label)
			style_header(c, header_fill)
			continue

		ws.merge_cells(start_row=h1, start_column=col, end_row=h3, end_column=col)
		c = ws.cell(row=h1, column=col, value=label)
		if label in YELLOW_HEADERS:
			style_header(c, yellow_fill)
		elif label in GREEN_HEADERS:
			style_header(c, green_fill)
		else:
			style_header(c)

	# Data
	row_num = data_start
	for sn, record in enumerate(records, start=1):
		ws.cell(row=row_num, column=1, value=sn).border = border
		for cidx, (field, col_label, is_num) in enumerate(DATA_COLUMNS, start=2):
			val = record.get(field) if field else None
			if field == "date_of_joining" and val:
				val = getdate(val).strftime("%d-%m-%Y")
			elif is_num and val is not None:
				val = round(float(val), 2)
			cell = ws.cell(row=row_num, column=cidx, value=val)
			cell.border = border
			if col_label in YELLOW_HEADERS:
				cell.fill = yellow_fill
			elif col_label in GREEN_HEADERS:
				cell.fill = green_fill
		row_num += 1

	current_totals = _sum_numeric_columns(records)
	if len(records) > 1:
		_write_total_row(ws, row_num, "S/TOTAL1", current_totals, border, Font(bold=True, size=9))
		row_num += 1

	_write_total_row(
		ws, row_num, f"GRAND TOTAL_{month.strftime('%B').upper()}", current_totals, border, Font(bold=True, size=9)
	)
	row_num += 1

	if prev_records:
		prev_totals = _sum_numeric_columns(prev_records)
		_write_total_row(
			ws,
			row_num,
			f"GRAND TOTAL_{prev_month.strftime('%B').upper()}",
			prev_totals,
			border,
			Font(bold=True, size=9),
		)
		row_num += 1
		variance = {k: current_totals.get(k, 0) - prev_totals.get(k, 0) for k in current_totals}
		_write_total_row(
			ws, row_num, "Variance", variance, border, Font(bold=True, size=9, color="FF0000")
		)
		row_num += 1

	footer = get_payroll_sheet_footer(
		company, payroll_month=month, payroll_excel_upload=payroll_excel_upload
	)
	row_num += 2
	ws.merge_cells(start_row=row_num, start_column=2, end_row=row_num, end_column=last_col)
	ws.cell(row=row_num, column=2, value=footer["date_line"]).font = Font(size=10)

	row_num += 2
	sig_cols = [2, max(2, last_col // 3), max(2, (2 * last_col) // 3)]
	for idx, sig in enumerate(footer["signatures"][:3]):
		col = sig_cols[idx] if idx < len(sig_cols) else 2 + idx * 8
		cell = ws.cell(row=row_num, column=col, value=sig["heading"])
		cell.font = Font(bold=True, size=9)
		ws.cell(row=row_num + 1, column=col, value=sig["name"]).font = Font(bold=True, size=9)
		ws.cell(row=row_num + 2, column=col, value=sig["title"]).font = Font(size=9)

	ws.column_dimensions["A"].width = 5
	for i in range(2, last_col + 1):
		ws.column_dimensions[get_column_letter(i)].width = 11

	buffer = io.BytesIO()
	wb.save(buffer)
	buffer.seek(0)
	return buffer.getvalue()


def _fetch_records(month, company: str | None):
	filters: dict = {"docstatus": 1, "payroll_month": month}
	if company:
		filters["company"] = company
	fields = list({col[0] for col in DATA_COLUMNS if col[0]})
	fields.append("payroll_period_label")
	return frappe.get_all(
		"Imported Payroll Record",
		filters=filters,
		fields=fields,
		order_by="employee_id_number asc",
	)


def _sum_numeric_columns(records: list[dict]) -> dict[int, float]:
	totals: dict[int, float] = {}
	for cidx, (field, _label, is_num) in enumerate(DATA_COLUMNS, start=2):
		if not is_num or not field:
			continue
		total = 0.0
		for record in records:
			try:
				total += float(record.get(field) or 0)
			except (TypeError, ValueError):
				pass
		totals[cidx] = total
	return totals


def _write_total_row(ws, row_num, label, totals, border, font):
	from openpyxl.styles import PatternFill

	header_fill = PatternFill("solid", fgColor="FDE9D9")
	yellow_fill = PatternFill("solid", fgColor="FFF2CC")
	green_fill = PatternFill("solid", fgColor="C6E0B4")
	ws.cell(row=row_num, column=1, value=label).font = font
	ws.cell(row=row_num, column=1).border = border
	for cidx, (_field, col_label, is_num) in enumerate(DATA_COLUMNS, start=2):
		cell = ws.cell(row=row_num, column=cidx)
		cell.border = border
		cell.font = font
		if is_num and cidx in totals:
			cell.value = round(totals[cidx], 2)
		if col_label in YELLOW_HEADERS:
			cell.fill = yellow_fill
		elif col_label in GREEN_HEADERS:
			cell.fill = green_fill
		else:
			cell.fill = header_fill
