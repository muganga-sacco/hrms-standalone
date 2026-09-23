from openpyxl import load_workbook

path = "/home/wngelique/frappe-bench/sites/mysite.local/private/files/Payroll in HRS.xlsx"
wb = load_workbook(path, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))[:8]
wb.close()
for i, row in enumerate(rows):
	print(i, list(row))
