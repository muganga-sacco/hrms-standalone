from openpyxl import load_workbook

p = "/home/wngelique/frappe-bench/sites/mysite.local/private/files/Payroll in HRSe361e0.xlsx"
wb = load_workbook(p, read_only=True, data_only=True)
rows = list(wb.active.iter_rows(values_only=True))
wb.close()
h = rows[7]
for i, c in enumerate(h):
    if c:
        print(i, repr(c))
print("--- row 10 data ---")
for i, c in enumerate(rows[10]):
    if c is not None and str(c).strip():
        print(i, repr(c))
