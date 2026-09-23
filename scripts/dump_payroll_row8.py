from openpyxl import load_workbook

path = "/home/wngelique/frappe-bench/sites/mysite.local/private/files/Payroll in HRS.xlsx"
wb = load_workbook(path, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
wb.close()
print("row8", list(rows[8])[:35])
print("row9", list(rows[9])[:35])
for r in rows[9:20]:
	if r and len(r) > 1 and r[1] and "003" in str(r[1]):
		print("data", list(r)[:35])
