from openpyxl import load_workbook

path = "/home/wngelique/frappe-bench/sites/mysite.local/private/files/Payroll in HRS.xlsx"
wb = load_workbook(path, read_only=True, data_only=True)
rows = list(wb.active.iter_rows(values_only=True))
wb.close()
for i in (17, 27, 28, 29):
    r = rows[i]
    for j, c in enumerate(r):
        if c is not None and str(c).strip():
            print(i, j, repr(c))
